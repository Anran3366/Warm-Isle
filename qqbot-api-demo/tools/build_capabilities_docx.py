from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
MD_PATH = ROOT / "docs" / "qqbot-official-capabilities.md"
DOCX_PATH = ROOT / "docs" / "qqbot-official-capabilities.docx"

NAVY = "132238"
INK = "1E2A38"
MUTED = "6C7888"
BLUE = "087EA4"
BLUE_DARK = "075E7D"
BLUE_LIGHT = "EAF6FA"
CORAL = "E8664A"
CORAL_LIGHT = "FDF0EC"
MINT = "16977F"
MINT_LIGHT = "EBF8F4"
LILAC = "6B5BD2"
LILAC_LIGHT = "F1EFFE"
PAPER = "F6F4EF"
LINE = "DCE5EA"
WHITE = "FFFFFF"
GRAY_BG = "F3F5F6"

FONT_ZH = "PingFang SC"
FONT_EN = "Aptos"
MONO = "SF Mono"

BODY_SIZE = Pt(10)
TABLE_SIZE = Pt(8.8)
CODE_SIZE = Pt(8.5)


def set_run_font(run, *, size=None, bold=None, color=None, name=FONT_EN, eastasia=FONT_ZH):
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), eastasia)


def shade(el, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    el.append(shd)


def set_cell_shading(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LINE, size=4):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_width(table, widths):
    table.autofit = False
    tblPr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = width


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    cant.set(qn("w:val"), "true")
    trPr.append(cant)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    set_run_font(run, size=Pt(8), color=MUTED)


def add_hyperlink(paragraph, url, text, color=BLUE, underline=True, size=None):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rPr.append(u)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rPr.append(c)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), FONT_EN)
    fonts.set(qn("w:hAnsi"), FONT_EN)
    fonts.set(qn("w:eastAsia"), FONT_ZH)
    rPr.append(fonts)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size.pt * 2)))
        rPr.append(sz)
    new_run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


TOKEN_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|https?://[^\s<]+)")


def add_inline(paragraph, text, *, size=BODY_SIZE, color=INK, bold=False):
    text = text.replace("\\|", "|")
    for part in TOKEN_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=Pt(max(size.pt - 0.6, 7.5)), color=BLUE_DARK, name=MONO, eastasia=FONT_ZH)
            shade(run._r.get_or_add_rPr(), "EDF4F7")
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = part[1:-1].split("](", 1)
            add_hyperlink(paragraph, url, label, size=size)
        elif part.startswith("http://") or part.startswith("https://"):
            add_hyperlink(paragraph, part, part, size=size)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def add_rule(paragraph, color=LINE, size=8, space=1):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), str(space))
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_callout(doc, lines, *, fill=BLUE_LIGHT, border=BLUE, label=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color=fill, size=2)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, start=180, bottom=120, end=160)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "18" if edge == "left" else "2")
        el.set(qn("w:color"), border if edge == "left" else fill)
        borders.append(el)
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.4
    if label:
        r = p.add_run(label + "  ")
        set_run_font(r, size=Pt(8), bold=True, color=border)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        add_inline(p, line, size=Pt(9.3), color=INK)
    return table


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table, color="E5EAEE", size=2)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F6F8F9")
    set_cell_margins(cell, top=90, start=140, bottom=90, end=120)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, size=CODE_SIZE, color="33414F", name=MONO, eastasia=FONT_ZH)


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pPr = p._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0" if level == 1 else "1")
    pPr.append(outline)
    if level == 1:
        pf.space_before = Pt(14)
        pf.space_after = Pt(7)
        pf.keep_with_next = True
        run = p.add_run(text)
        set_run_font(run, size=Pt(17), bold=True, color=BLUE_DARK)
        add_rule(p, color="9ED6E6", size=8, space=4)
    elif level == 2:
        pf.space_before = Pt(10)
        pf.space_after = Pt(4)
        pf.keep_with_next = True
        run = p.add_run(text)
        set_run_font(run, size=Pt(12), bold=True, color=INK)
    else:
        pf.space_before = Pt(8)
        pf.space_after = Pt(3)
        pf.keep_with_next = True
        run = p.add_run(text)
        set_run_font(run, size=Pt(10.5), bold=True, color=BLUE_DARK)
    return p


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_borders(table)
    if cols == 2:
        widths = [Cm(4.2), Cm(12.8)]
    elif cols == 3:
        widths = [Cm(3.0), Cm(4.3), Cm(9.7)]
    elif cols == 4:
        widths = [Cm(3.1), Cm(3.1), Cm(3.1), Cm(7.7)]
    elif cols == 5:
        widths = [Cm(2.5), Cm(2.4), Cm(2.4), Cm(3.4), Cm(6.3)]
    else:
        widths = [Cm(17.0 / cols)] * cols
    set_table_width(table, widths)

    for ridx, row_data in enumerate(rows):
        row = table.rows[ridx]
        prevent_row_split(row)
        if ridx == 0:
            repeat_header(row)
        for cidx in range(cols):
            cell = row.cells[cidx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if ridx == 0:
                set_cell_shading(cell, NAVY)
            elif ridx % 2 == 0:
                set_cell_shading(cell, "F8FAFB")
            text = row_data[cidx] if cidx < len(row_data) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.25
            add_inline(
                p,
                text,
                size=TABLE_SIZE,
                color=WHITE if ridx == 0 else INK,
                bold=(ridx == 0),
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_table(lines, idx):
    rows = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        line = lines[idx].strip()
        if not re.match(r"^\|[\s:\-\|]+\|$", line):
            cells = [c.strip() for c in line.strip("|").split("|")]
            rows.append(cells)
        idx += 1
    if rows:
        add_table(DOC, rows)
    return idx


def cover(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)
    section.different_first_page_header_footer = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run("QQ BOT  /  OPEN PLATFORM  /  API V2")
    set_run_font(r, size=Pt(8.5), bold=True, color=BLUE)

    band = doc.add_table(rows=1, cols=1)
    set_table_borders(band, color=NAVY, size=2)
    cell = band.cell(0, 0)
    set_cell_shading(cell, NAVY)
    set_cell_margins(cell, top=380, start=300, bottom=330, end=300)
    cp = cell.paragraphs[0]
    cp.paragraph_format.space_after = Pt(0)
    r = cp.add_run("QQ 官方机器人\n能力地图")
    set_run_font(r, size=Pt(34), bold=True, color=WHITE)
    r = cp.add_run("\n\n从消息、事件到群 / 频道管理与 Agent 接入")
    set_run_font(r, size=Pt(12), color="CDE7F0")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("开发者视角的完整能力清单、限额与边界")
    set_run_font(r, size=Pt(18), bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("不是个人号自动化，而是平台内的官方机器人服务")
    set_run_font(r, size=Pt(10.5), color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    add_rule(p, color=CORAL, size=18, space=2)

    meta = doc.add_table(rows=1, cols=3)
    set_table_borders(meta, color=WHITE, size=2)
    stats = [
        ("04", "会话场景", "单聊 · 群聊 · 频道 · 私信"),
        ("16", "能力章节", "消息到 Agent 接入"),
        ("1000", "主动消息 / 日", "单用户上限（当前文档）"),
    ]
    for i, (num, label, desc) in enumerate(stats):
        cell = meta.cell(0, i)
        set_cell_shading(cell, [BLUE_LIGHT, CORAL_LIGHT, MINT_LIGHT][i])
        set_cell_margins(cell, top=160, start=160, bottom=150, end=160)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(num)
        set_run_font(r, size=Pt(24), bold=True, color=[BLUE, CORAL, MINT][i])
        r = p.add_run("\n" + label)
        set_run_font(r, size=Pt(9), bold=True, color=INK)
        r = p.add_run("\n" + desc)
        set_run_font(r, size=Pt(7.5), color=MUTED)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(34)
    r = p.add_run("资料基线  2026-10-21 官方文档\n访问日期  2026-10-26\n整理  DSH Agent")
    set_run_font(r, size=Pt(8.5), color=MUTED)
    p.paragraph_format.line_spacing = 1.7

    doc.add_page_break()


def toc(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("本文速览")
    set_run_font(r, size=Pt(20), bold=True, color=INK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("先看结论，再看能力矩阵；第五章之后是群 / 频道 / Agent 的专项能力。")
    set_run_font(r, size=Pt(9.5), color=MUTED)

    sections = [
        ("00", "一句话结论", "官方机器人是什么、不是什么"),
        ("01", "能力总览", "12 类能力一页看完"),
        ("02", "四种会话场景", "单聊、群聊、子频道、频道私信"),
        ("03", "消息能力", "文本、Markdown、富媒体、按钮、回复与召回"),
        ("04", "事件与实时通信", "WebSocket / Webhook、Intents 与事件清单"),
        ("05", "机器人与运营入口", "资料、分享链接、菜单、指令面板"),
        ("06", "群聊管理能力", "群信息、审批、禁言、黑名单、成员"),
        ("07", "频道（Guild）能力", "子频道、私信、表态、审核"),
        ("08", "Agent / 大模型接入", "流式输出、扫码绑定、OpenClaw / DSH"),
        ("09", "鉴权、权限与合规", "Token、白名单、运营规范"),
        ("10", "官方机器人不能做什么", "边界与不可行清单"),
        ("11", "典型产品形态", "7 种可落地玩法"),
        ("12", "配额与限制速查", "一张表看全部关键数字"),
        ("13", "文档口径与踩坑提醒", "旧文档 vs 当前规则"),
        ("14", "给当前项目的落地建议", "被动优先、主动克制"),
        ("15", "官方文档索引", "主要来源链接"),
    ]
    table = doc.add_table(rows=len(sections), cols=3)
    set_table_borders(table, color=LINE, size=2)
    set_table_width(table, [Cm(1.2), Cm(5.0), Cm(10.8)])
    for idx, (num, title, desc) in enumerate(sections):
        row = table.rows[idx]
        prevent_row_split(row)
        cells = row.cells
        for c in cells:
            set_cell_margins(c, top=75, start=70, bottom=75, end=90)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if idx % 2 == 0:
            for c in cells:
                set_cell_shading(c, "F8FAFB")
        p = cells[0].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(num)
        set_run_font(r, size=Pt(9), bold=True, color=BLUE_DARK, name=MONO)
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, title, size=Pt(9.5), color=INK, bold=True)
        p = cells[2].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, desc, size=Pt(8.5), color=MUTED)

    doc.add_paragraph()
    add_callout(
        doc,
        [
            "**先记住三件事：** ① 所有主动触达都要先有 openid / group_openid；② 用户/群可以关闭主动消息，关了就会失败；"
            "③ 群管理、自定义按钮、频道 Markdown 等能力带白名单 / 内邀门槛，不能默认全量可用。",
        ],
        fill=CORAL_LIGHT,
        border=CORAL,
        label="阅读提示",
    )
    doc.add_page_break()


def header_footer(doc):
    section = doc.sections[0]
    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("QQ 官方机器人能力地图  ·  2026-10")
    set_run_font(r, size=Pt(7.5), color="9AA6B2")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("QQ BOT API V2  /  CAPABILITY MAP     ")
    set_run_font(r, size=Pt(7.5), color="9AA6B2")
    add_page_field(p)


def body(doc):
    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    # 跳过标题和开头的说明块，从“0. 一句话结论”开始
    start = 0
    for i, line in enumerate(lines):
        if line.startswith("## 0."):
            start = i
            break
    lines = lines[start:]
    idx = 0
    code_mode = False
    code_lines = []
    quote_lines = []

    def flush_quote():
        nonlocal quote_lines
        if quote_lines:
            add_callout(doc, quote_lines, fill=BLUE_LIGHT, border=BLUE, label="")
            quote_lines = []

    while idx < len(lines):
        raw = lines[idx]
        line = raw.strip()

        if code_mode:
            if line.startswith("```"):
                add_code_block(doc, code_lines)
                code_lines = []
                code_mode = False
            else:
                code_lines.append(raw.rstrip())
            idx += 1
            continue

        if line.startswith("```"):
            flush_quote()
            code_mode = True
            idx += 1
            continue

        if line.startswith("|"):
            flush_quote()
            idx = parse_table(lines, idx)
            continue

        if line.startswith(">"):
            quote_lines.append(line[1:].strip())
            idx += 1
            continue
        flush_quote()

        if not line:
            idx += 1
            continue

        if line == "---":
            idx += 1
            continue

        if line.startswith("### "):
            add_heading(doc, line[4:], 2)
            idx += 1
            continue
        if line.startswith("## "):
            text = line[3:]
            p = add_heading(doc, text, 1)
            if text.startswith("3."):
                p.paragraph_format.page_break_before = True
            idx += 1
            continue
        if line.startswith("# "):
            idx += 1
            continue

        match_num = re.match(r"^(\d+)\.\s+", line)
        if match_num:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.66)
            p.paragraph_format.first_line_indent = Cm(-0.46)
            p.paragraph_format.space_after = Pt(2.5)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run(match_num.group(1) + ".  ")
            set_run_font(r, size=BODY_SIZE, bold=True, color=BLUE_DARK)
            add_inline(p, line[match_num.end():])
            idx += 1
            continue

        if line.startswith("- ") or raw.startswith("  - "):
            nested = raw.startswith("  - ")
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.08 if nested else 0.62)
            p.paragraph_format.first_line_indent = Cm(-0.32)
            p.paragraph_format.space_after = Pt(2.2)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run("– " if nested else "• ")
            set_run_font(r, size=Pt(9.5), bold=False, color=MUTED if nested else BLUE_DARK)
            add_inline(p, line[2:])
            idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.line_spacing = 1.55
        add_inline(p, line)
        idx += 1

    if code_mode and code_lines:
        add_code_block(doc, code_lines)


def finalize(doc):
    # 默认样式兜底
    normal = doc.styles["Normal"]
    normal.font.name = FONT_EN
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_ZH)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.5

    # 页脚段落字体
    for section in doc.sections:
        for p in section.footer.paragraphs:
            for run in p.runs:
                if run.font.size is None:
                    set_run_font(run, size=Pt(7.5), color=MUTED)

    core = doc.core_properties
    core.title = "QQ 官方机器人能力地图"
    core.subject = "QQ 机器人开放平台 API v2 能力、限额与边界"
    core.author = "DSH Agent"
    core.keywords = "QQ Bot, API v2, 官方机器人, 群管理, 频道, Agent"
    core.comments = "资料基线 2026-10-21；访问日期 2026-10-26"
    core.language = "zh-CN"

    doc.save(DOCX_PATH)


def main():
    global DOC
    DOC = Document()
    cover(DOC)
    toc(DOC)
    body(DOC)
    header_footer(DOC)
    finalize(DOC)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
